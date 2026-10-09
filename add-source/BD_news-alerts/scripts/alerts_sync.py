"""Compare the Google Alerts sheet with the alerts connected in sources.toml.

Read-only: prints what to add; it never edits sources.toml or the sheet.

    uv run python .claude/skills/news-alerts/scripts/alerts_sync.py            # compare only
    uv run python .claude/skills/news-alerts/scripts/alerts_sync.py --check    # also fetch each NEW feed once

For every sheet row it reports one of:
  connected   the link is already in sources.toml
  NEW         a valid feed link not yet connected (prints a ready TOML line)
  NOT A LINK  something else was pasted in the link column (e.g. the query text)
  no link     nothing pasted yet
and lists connected alerts whose link is not in the sheet.

--check fetches each NEW link once with the app's polite client and confirms it is an
Atom feed whose title matches the row's query, which catches a link pasted into the
wrong row. Only feed URLs (google.com/alerts/feeds/...) are ever fetched.
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

from bdleads import config

FEED_RE = re.compile(r"^https://www\.google\.com/alerts/feeds/\d+/\d+$")
TITLE_RE = re.compile(r"<title[^>]*>(.*?)</title>", re.S)


def read_sheet(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Alerts"] if "Alerts" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "") for h in next(rows)]

    def col(word: str) -> int:
        for i, h in enumerate(header):
            if word.lower() in h.lower():
                return i
        sys.exit(f"Column containing '{word}' not found in {path}; header is {header}")

    i_no, i_q, i_link = col("#"), col("query"), col("RSS")
    out = []
    for r in rows:
        if not r or r[i_q] in (None, ""):
            continue
        out.append({"row": r[i_no], "query": str(r[i_q]).strip(),
                    "link": str(r[i_link] or "").strip()})
    wb.close()
    return out


def short_name(query: str) -> str:
    q = query.replace("site:linkedin.com", "LinkedIn -").replace('"', "")
    q = re.sub(r"\s+", " ", q).strip()
    return "Google Alert: " + (q[:60].rstrip() + ("..." if len(q) > 60 else ""))


def check_feed(url: str, query: str) -> str:
    from bdleads.sources import PoliteClient

    http = PoliteClient()
    try:
        r = http.get(url)
    except Exception as e:  # noqa: BLE001
        return f"FAILED to fetch: {e}"
    finally:
        http.close()
    body = r.text
    if "<feed" not in body[:500]:
        return f"NOT A FEED (HTTP {r.status_code})"
    m = TITLE_RE.search(body)
    title = (m.group(1) if m else "").replace("&quot;", '"').replace("&amp;", "&").strip()
    entries = body.count("<entry>")
    norm = lambda s: re.sub(r"\s+", " ", s.replace('"', "")).strip().lower()  # noqa: E731
    match = norm(title).endswith(norm(query))
    verdict = "title matches the row's query" if match else f"TITLE DOES NOT MATCH THE ROW: {title!r}"
    return f"OK feed, {entries} entries now; {verdict}"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", default="docs/Google_Alerts_List.xlsx")
    ap.add_argument("--check", action="store_true", help="fetch each NEW feed once to verify it")
    a = ap.parse_args()

    sheet = read_sheet(Path(a.sheet))
    connected = {f["url"]: f for f in config.sources()["rss"]["feeds"] if f.get("google_alert")}
    today = date.today().isoformat()

    counts: dict[str, int] = {}
    new_lines = []
    print(f"Sheet: {a.sheet} ({len(sheet)} rows); connected alerts in sources.toml: {len(connected)}\n")
    for s in sheet:
        link = s["link"]
        if not link:
            status = "no link"
        elif link in connected:
            status = "connected"
        elif FEED_RE.match(link):
            status = "NEW"
        else:
            status = "NOT A LINK"
        counts[status] = counts.get(status, 0) + 1
        if status == "no link":
            continue
        line = f"row {s['row']:>2}  {status:<10}  {s['query']}"
        if status == "connected":
            line += f"\n          as: {connected[link]['name']}"
        if status == "NOT A LINK":
            line += f"\n          link column contains: {link[:100]!r}"
        if status == "NEW":
            name = short_name(s["query"])
            toml = f'  {{ name = "{name}", google_alert = true, added = "{today}", url = "{link}" }},'
            new_lines.append(toml)
            line += f"\n          {toml.strip()}"
            if a.check:
                line += f"\n          check: {check_feed(link, s['query'])}"
        print(line)

    sheet_links = {s["link"] for s in sheet}
    orphans = [f for u, f in connected.items() if u not in sheet_links]
    print("\nSummary: " + ", ".join(f"{k}: {v}" for k, v in sorted(counts.items())))
    if orphans:
        print("Connected but not in the sheet (deleted from the sheet, or added another way):")
        for f in orphans:
            print(f"  {f['name']}  {f['url']}")
    if new_lines:
        print("\nTOML lines to insert at the end of the Google Alerts block in [rss] feeds"
              " (use the Edit tool; check names are unique):")
        print("\n".join(new_lines))


if __name__ == "__main__":
    main()
