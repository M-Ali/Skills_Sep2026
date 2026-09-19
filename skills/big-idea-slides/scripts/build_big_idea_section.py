"""Build the Big Idea section from a JSON spec.

    uv run --with python-pptx python build_big_idea_section.py spec.json \
        [--base deck.pptx --insert-at N]

Slides: divider, consumer insight, category insight, brand insight, convergence,
strategy + platform, and (if `commitments` is present) the price of the idea.

- Without --base: a new 16:9 deck holding only the section.
- With --base: a copy of that deck with the section inserted at --insert-at (0-based), or
  appended when --insert-at is omitted.
- Never overwrites: writes the first free name among out.pptx, out-v2.pptx, ...

Checks that stop a build (things that would mislead a client), and warnings that don't.
Spec shape: ../assets/example_spec.json. Layout: ../references/slide-specs.md.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x10, 0x2A, 0x43)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x74, 0x80)
ACCENT = RGBColor(0xC8, 0x10, 0x2E)
PALE = RGBColor(0xF2, 0xF5, 0xF8)
DIM = RGBColor(0x9F, 0xB4, 0xC7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"

ROLE = {"consumer": "Relevant", "category": "Distinctive", "brand": "Ownable"}
STEP_SOURCE = {"consumer": "FROM THE CONSUMER PULSE",
               "category": "FROM COMPETITORS' OWNED CHANNELS",
               "brand": "FROM THE CLIENT DOCUMENTS"}


# ---------- primitives ----------

def rect(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, paras, size=12, bold=False, color=INK, space=6,
         align=PP_ALIGN.LEFT, italic=False):
    """paras: str, or list of str / (label, body). A label renders as a bold run-in."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        p.alignment = align
        label, body = para if isinstance(para, tuple) else (None, para)
        if label:
            r = p.add_run()
            r.text = label + "  "
            _font(r, size, True, color, italic)
        r = p.add_run()
        r.text = body
        _font(r, size, bold, color, italic)
    return tb


def _font(run, size, bold, color, italic=False):
    run.font.name, run.font.size = FONT, Pt(size)
    run.font.bold, run.font.italic = bold, italic
    run.font.color.rgb = color


def header(prs, s, kicker, title):
    rect(s, 0, 0, prs.slide_width / 914400, 1.0, NAVY)
    text(s, 0.6, 0.14, 12.0, 0.3, kicker.upper(), 10, True, DIM, space=0)
    text(s, 0.6, 0.42, 12.1, 0.5, title, 22, True, WHITE, space=0)


# ---------- validation ----------

def check(spec):
    errors, warnings = [], []
    insights = spec["insights"]
    for kind in ("consumer", "category", "brand"):
        ins = insights.get(kind)
        if not ins:
            errors.append(f"missing the {kind} insight - an idea on two insights is not an idea")
            continue
        for field in ("line", "short", "source"):
            if not ins.get(field):
                errors.append(f"{kind} insight has no `{field}`"
                              + (" (the convergence slide is built from it)" if field == "short" else ""))
        ev = ins.get("evidence") or []
        if not ev:
            errors.append(f"{kind} insight has no evidence")
        elif len(ev) < 2:
            warnings.append(f"{kind} insight rests on a single piece of evidence")
        elif len(ev) > 5:
            warnings.append(f"{kind} insight has {len(ev)} evidence bullets; 5 fit the slide")
        line = ins.get("line", "")
        if ";" in line or len(re.findall(r"[.!?](?:\s|$)", line.rstrip())) > 1:
            warnings.append(f"{kind} insight line holds more than one thought - split it and choose")

    idea = spec["idea"]
    norm = lambda s: re.sub(r"[^a-z0-9 ]", "", s.lower()).strip()
    if norm(idea["line"]) == norm(spec["strategy"]["platform"]):
        errors.append("the Big Idea and the comms platform are the same words - "
                      "the idea is a thought, the platform is the briefing line")
    for field in ("meaning", "test"):
        if not idea.get(field):
            errors.append(f"idea has no `{field}`")

    st = spec["strategy"]
    missing = [w for w in ("from", "to", "by") if f" {w} " not in f" {st['statement'].lower()} "]
    if missing:
        warnings.append(f"strategy statement has no {' / '.join(missing)} - "
                        "'Re-frame X from ... to ... by ...' is the shape that briefs people")
    alts = st.get("alternatives") or []
    if len(alts) < 2:
        errors.append("fewer than two alternatives considered - without the road not taken this "
                      "reads as the only idea anyone had")
    if not any(a.get("recommended") for a in alts):
        errors.append("no alternative is marked `recommended`")
    if len(alts) > 4:
        warnings.append(f"{len(alts)} alternatives; 4 entries fit the panel")

    if spec.get("requires_commitments") and not spec.get("commitments", {}).get("asks"):
        errors.append("requires_commitments is true but no commitments are listed - "
                      "an idea whose proof never arrives damages the brand")

    if errors:
        raise SystemExit("Refusing to build:\n  - " + "\n  - ".join(errors))
    for w in warnings:
        print(f"warning: {w}")


# ---------- slides ----------

def divider(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, prs.slide_width / 914400, prs.slide_height / 914400, NAVY)
    rect(s, 0.9, 2.3, 1.6, 0.06, ACCENT)
    text(s, 0.9, 2.6, 11.4, 1.0, d.get("title", "The Big Idea"), 40, True, WHITE)
    text(s, 0.9, 3.55, 11.4, 1.2,
         [d.get("subtitle", "Consumer insight  +  Category insight  +  Brand insight"),
          d.get("subtitle2", "Three single lines, combined into one organising thought.")],
         16, False, DIM)


def insight_slide(prs, step, kind, ins):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(prs, s, f"STEP {step}  ·  {kind.upper()} INSIGHT  ·  {ins.get('from', STEP_SOURCE[kind])}",
           ins["title"])
    rect(s, 0.6, 1.35, 0.08, 1.75, ACCENT)
    text(s, 0.95, 1.35, 11.7, 1.75, f"“{ins['line']}”", 23, True, NAVY, space=0)
    if ins.get("tag"):
        rect(s, 10.9, 1.05, 1.8, 0.26, ACCENT)
        text(s, 10.9, 1.08, 1.8, 0.22, ins["tag"].upper(), 9, True, WHITE, space=0, align=PP_ALIGN.CENTER)
    text(s, 0.6, 3.4, 6.9, 0.3, "EVIDENCE", 10, True, MUTED, space=0)
    text(s, 0.6, 3.72, 6.9, 3.1, [f"·  {e}" for e in ins["evidence"]], 11, False, INK, space=7)
    rect(s, 7.9, 3.4, 4.8, 3.4, PALE)
    text(s, 8.15, 3.55, 4.35, 0.3, ins["panel_title"].upper(), 10, True, MUTED, space=0)
    text(s, 8.15, 3.9, 4.35, 2.85, ins["panel"], 10.5, False, NAVY)
    text(s, 0.6, 7.05, 12.1, 0.3, ins["source"], 8.5, False, MUTED, space=0)


def convergence(prs, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(prs, s, "STEP 4  ·  THE BIG IDEA", spec["idea"].get("title", "Three insights, one organising thought"))
    for i, kind in enumerate(("consumer", "category", "brand")):
        x = 0.6 + i * 4.1
        rect(s, x, 1.35, 3.9, 0.05, ACCENT)
        text(s, x, 1.55, 3.9, 0.3, f"{kind.upper()} INSIGHT  ·  {ROLE[kind]}", 10, True, MUTED, space=0)
        text(s, x, 1.9, 3.9, 1.4, spec["insights"][kind]["short"], 14, True, NAVY, space=0)
    text(s, 0.6, 3.3, 12.1, 0.5, "▼", 22, True, ACCENT, space=0, align=PP_ALIGN.CENTER)
    rect(s, 0.6, 3.85, 12.1, 1.55, NAVY)
    text(s, 0.9, 3.98, 11.5, 0.3, "THE BIG IDEA", 10, True, DIM, space=0)
    line = spec["idea"]["line"]
    text(s, 0.9, 4.3, 11.5, 1.0, line, 28 if len(line) > 60 else 32, True, WHITE, space=0)
    text(s, 0.6, 5.65, 12.1, 0.9, spec["idea"]["meaning"], 14, False, INK, space=0)
    text(s, 0.6, 6.6, 12.1, 0.6, [("Test:", spec["idea"]["test"])], 10.5, False, MUTED, space=0)


def strategy_slide(prs, spec):
    st, brand = spec["strategy"], spec["brand"]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(prs, s, "FROM IDEA TO STRATEGY", st.get("title", "Strategy statement and comms platform"))
    text(s, 0.6, 1.35, 7.2, 0.3, "STRATEGY", 10, True, MUTED, space=0)
    text(s, 0.6, 1.65, 7.2, 1.9, st["statement"], 16, True, NAVY, space=0)
    rect(s, 0.6, 3.7, 7.2, 1.3, PALE)
    text(s, 0.85, 3.82, 6.8, 0.3, "COMMS PLATFORM", 10, True, MUTED, space=0)
    text(s, 0.85, 4.12, 6.8, 0.8, [(f"Position {brand} as", st["platform"])], 19, True, ACCENT, space=0)
    text(s, 0.6, 5.25, 7.2, 0.3, "WHY IT HOLDS", 10, True, MUTED, space=0)
    text(s, 0.6, 5.55, 7.2, 1.6,
         [("Relevant:", st["why"]["relevant"]), ("Distinctive:", st["why"]["distinctive"]),
          ("Ownable:", st["why"]["ownable"])], 11, False, INK, space=5)
    rect(s, 8.2, 1.35, 4.5, 5.8, PALE)
    text(s, 8.45, 1.5, 4.05, 0.3, "ALTERNATIVES CONSIDERED", 10, True, MUTED, space=0)
    entries = [(("Recommended - " if a.get("recommended") else "") + a["name"], a["tradeoff"])
               for a in st["alternatives"]]
    if st.get("swap_test"):
        entries.append(("Swap test:", st["swap_test"]))
    text(s, 8.45, 1.85, 4.05, 5.2, entries, 10.5, False, NAVY, space=9)


def price_slide(prs, spec):
    c = spec["commitments"]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(prs, s, "THE PRICE OF THE IDEA", c.get("title", "What it asks, and how it lives in media"))
    text(s, 0.6, 1.35, 6.0, 0.3, c.get("asks_title", "What the client must commit to").upper(),
         10, True, MUTED, space=0)
    y = 1.75
    for i, a in enumerate(c["asks"], 1):
        rect(s, 0.6, y, 0.05, 0.72, ACCENT)
        text(s, 0.8, y + 0.02, 5.8, 0.3, f"{i}  {a['name']}", 13, True, NAVY, space=0)
        text(s, 0.8, y + 0.35, 5.8, 0.4, a["detail"], 11, False, INK, space=0)
        y += 0.88
    if c.get("media_principles"):
        text(s, 7.0, 1.35, 5.7, 0.3, "MEDIA PRINCIPLES", 10, True, MUTED, space=0)
        top = 1.7
        if c.get("channel_insight"):
            text(s, 7.0, 1.7, 5.7, 0.9, f"“{c['channel_insight']}”", 13, True, NAVY,
                 space=0, italic=True)
            top = 2.65
        text(s, 7.0, top, 5.7, 3.6, [(p["name"], p["detail"]) for p in c["media_principles"]],
             11.5, False, INK, space=9)
    if c.get("conclusion"):
        rect(s, 0.6, 6.35, 12.1, 0.8, NAVY)
        text(s, 0.85, 6.47, 11.6, 0.6, c["conclusion"], 12.5, True, WHITE, space=0)


# ---------- deck handling ----------

def free_path(out: Path) -> Path:
    if not out.exists():
        return out
    n = 2
    while (p := out.with_name(f"{out.stem}-v{n}{out.suffix}")).exists():
        n += 1
    return p


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", type=Path)
    ap.add_argument("--base", type=Path, help="existing deck to insert into (copied, never edited)")
    ap.add_argument("--insert-at", type=int, help="0-based slide index for the divider (default: end)")
    args = ap.parse_args()

    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    check(spec)

    if args.base:
        prs = Presentation(str(args.base))
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    before = len(prs.slides)

    divider(prs, spec.get("divider", {}))
    for step, kind in enumerate(("consumer", "category", "brand"), start=1):
        insight_slide(prs, step, kind, spec["insights"][kind])
    convergence(prs, spec)
    strategy_slide(prs, spec)
    if spec.get("commitments", {}).get("asks"):
        price_slide(prs, spec)

    if args.insert_at is not None:
        ids = prs.slides._sldIdLst
        new = list(ids)[before:]
        for el in new:
            ids.remove(el)
        for i, el in enumerate(new):
            ids.insert(args.insert_at + i, el)

    out = free_path(Path(spec["out"]) if Path(spec["out"]).is_absolute() else args.spec.parent / spec["out"])
    prs.save(out)
    print(f"{out}  ({len(prs.slides)} slides; section = {len(prs.slides) - before} slides)")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
